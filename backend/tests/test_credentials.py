"""The first password an administrator hands out can be looked at again - until the employee chooses their own.

Encrypted at rest, administrators only, audited every time, gone the moment the employee changes it."""

from __future__ import annotations

import io

from openpyxl import load_workbook

from app.core.timeutils import utcnow
from app.models.distribution import EmployeeCredential
from app.models.system import AuditLog
from tests.conftest import auth_headers, login

NEW = {"email": "new.hire@example.com", "full_name": "New Hire"}


def create(client, headers, **extra):
    answer = client.post("/api/v1/employees", headers=headers, json={**NEW, **extra})
    assert answer.status_code == 201, answer.text
    return answer.json()


def test_a_generated_password_can_be_seen_again_by_the_administrator(client, as_admin, db):
    created = create(client, as_admin)
    employee_id, password = created["employee"]["id"], created["temporary_password"]
    seen = client.get(f"/api/v1/employees/{employee_id}/credentials", headers=as_admin)
    assert seen.status_code == 200 and seen.headers["cache-control"] == "no-store"
    body = seen.json()
    assert body["available"] is True and body["password"] == password and body["kind"] == "generated"
    assert body["employee_code"] == created["employee"]["employee_code"] and body["must_change_password"] is True
    assert body["set_by"] and body["view_count"] == 1
    assert client.get(f"/api/v1/employees/{employee_id}/credentials", headers=as_admin).json()["view_count"] == 2
    assert db.query(AuditLog).filter(AuditLog.action == "employee.credentials_viewed").count() == 2  # every viewing is written down


def test_a_password_the_administrator_typed_can_be_seen_too(client, as_admin):
    created = create(client, as_admin, password="Typed-By-Boss-77")
    assert created["temporary_password"] is None  # it is not echoed back at creation
    body = client.get(f"/api/v1/employees/{created['employee']['id']}/credentials", headers=as_admin).json()
    assert body["available"] and body["password"] == "Typed-By-Boss-77" and body["kind"] == "admin_set"


def test_it_is_stored_encrypted_not_readable_and_not_decryptable_with_another_key(client, as_admin, db, monkeypatch):
    created = create(client, as_admin)
    row = db.get(EmployeeCredential, created["employee"]["id"])
    assert created["temporary_password"] not in row.ciphertext and row.ciphertext.startswith("gAAAA")  # a Fernet token

    from app.core.config import reset_settings_cache

    monkeypatch.setenv("JWT_SECRET", "another-secret-0123456789-0123456789-0123456789-xx")
    reset_settings_cache()
    try:
        from app.services import credential_vault

        assert credential_vault.read(db, created["employee"]["id"]) is None  # a different key shows nothing
    finally:
        monkeypatch.undo()
        reset_settings_cache()


def test_when_the_employee_chooses_their_own_password_nobody_can_see_it(client, as_admin):
    created = create(client, as_admin)
    employee_id = created["employee"]["id"]
    signed_in = login(client, NEW["email"], created["temporary_password"])
    headers = {"Authorization": f"Bearer {signed_in['access_token']}"}
    changed = client.post(
        "/api/v1/auth/change-password", headers=headers, json={"current_password": created["temporary_password"], "new_password": "My-Own-Secret-42x"}
    )
    assert changed.status_code == 200, changed.text
    body = client.get(f"/api/v1/employees/{employee_id}/credentials", headers=as_admin).json()
    assert body["available"] is False and body["password"] is None


def test_a_reset_shows_the_new_password_and_a_deactivated_account_shows_none(client, as_admin):
    created = create(client, as_admin)
    employee_id = created["employee"]["id"]
    reset = client.post(f"/api/v1/employees/{employee_id}/reset-password", headers=as_admin, json={}).json()
    assert client.get(f"/api/v1/employees/{employee_id}/credentials", headers=as_admin).json()["password"] == reset["temporary_password"]
    assert client.post(f"/api/v1/employees/{employee_id}/deactivate", headers=as_admin).status_code == 200
    assert client.get(f"/api/v1/employees/{employee_id}/credentials", headers=as_admin).json()["available"] is False


def test_a_password_kept_for_too_long_is_gone(client, as_admin, db):
    created = create(client, as_admin)
    row = db.get(EmployeeCredential, created["employee"]["id"])
    from datetime import timedelta

    row.created_at = utcnow() - timedelta(days=31)
    db.commit()
    assert client.get(f"/api/v1/employees/{created['employee']['id']}/credentials", headers=as_admin).json()["available"] is False
    from app.services import credential_vault

    assert credential_vault.purge_expired(db) == 1 and db.query(EmployeeCredential).count() == 0


def test_only_administrators_may_look(client, make, as_admin):
    created = create(client, as_admin)
    manager = make.employee(role="manager", code="MGRX", email="mgrx@example.com")
    worker = make.employee(code="WRKX", email="wrkx@example.com")
    for who in (manager, worker):
        headers = auth_headers(client, who)
        assert client.get(f"/api/v1/employees/{created['employee']['id']}/credentials", headers=headers).status_code == 403
        assert client.get("/api/v1/employees/credentials.xlsx", headers=headers).status_code == 403
    assert client.get(f"/api/v1/employees/{created['employee']['id']}/credentials").status_code == 401


def test_an_employee_who_does_not_exist_is_a_404(client, as_admin):
    assert client.get("/api/v1/employees/999999/credentials", headers=as_admin).status_code == 404


def test_bulk_creation_remembers_every_password(client, as_admin):
    batch = [{"email": f"bulk{i}@example.com", "full_name": f"Bulk Person {i}"} for i in range(5)]
    done = client.post("/api/v1/employees/bulk", headers=as_admin, json={"employees": batch}).json()
    assert done["created"] == 5
    for item in done["results"]:
        seen = client.get(f"/api/v1/employees/{item['employee']['id']}/credentials", headers=as_admin).json()
        assert seen["password"] == item["temporary_password"]


def test_the_login_sheet_lists_the_pending_logins_as_text(client, as_admin, db):
    first = create(client, as_admin)
    odd = create(client, as_admin, email="odd@example.com", full_name="Odd Password", password="=HYPERLINK-ish-9Zz")  # looks like a formula
    done = create(client, as_admin, email="done@example.com", full_name="Already Done")
    signed_in = login(client, "done@example.com", done["temporary_password"])
    client.post(
        "/api/v1/auth/change-password",
        headers={"Authorization": f"Bearer {signed_in['access_token']}"},
        json={"current_password": done["temporary_password"], "new_password": "Chosen-By-Them-55q"},
    )

    answer = client.get("/api/v1/employees/credentials.xlsx", headers=as_admin)
    assert answer.status_code == 200 and "spreadsheetml" in answer.headers["content-type"] and answer.headers["cache-control"] == "no-store"
    sheet = load_workbook(io.BytesIO(answer.content)).active
    rows = [[cell.value for cell in row] for row in sheet.iter_rows()]
    assert rows[0][0] == "Employee ID" and len(rows) == 3  # the header and the two employees who have not chosen yet
    by_email = {r[2]: r for r in rows[1:]}
    assert set(by_email) == {NEW["email"], "odd@example.com"}
    assert by_email[NEW["email"]][4] == first["temporary_password"]
    assert by_email["odd@example.com"][4] == "=HYPERLINK-ish-9Zz"
    assert sheet["E3"].data_type == "s" or sheet["E2"].data_type == "s"  # text, not a formula
    assert odd and db.query(AuditLog).filter(AuditLog.action == "employee.credentials_exported").count() == 1
    only_one = client.get(f"/api/v1/employees/credentials.xlsx?ids={first['employee']['id']}", headers=as_admin)
    assert len(list(load_workbook(io.BytesIO(only_one.content)).active.iter_rows())) == 2
    assert client.get("/api/v1/employees/credentials.xlsx?ids=1;drop", headers=as_admin).status_code == 422
