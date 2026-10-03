"""Reading a contact sheet safely, one row at a time: CSV and XLSX of any size, never the whole file in memory.

What is guarded against: a file that is not a sheet at all (binary data renamed to .csv), text in any of the usual encodings, a
delimiter that is a semicolon or a tab, cells of absurd size, too many columns, and the `.xlsx` "zip bomb" (a small file that
unpacks to gigabytes, or whose entries are compressed more than any real sheet) - checked BEFORE a byte is unpacked.
"""

from __future__ import annotations

import codecs
import csv
import io
import zipfile
from collections.abc import Iterator
from pathlib import Path
from typing import Any

from app.core.config import get_settings

_SAMPLE_BYTES = 64 * 1024
_MAX_ZIP_ENTRIES = 5000
_MAX_RATIO = 500  # a real sheet compresses 5-20 times; this only catches deliberate bombs


class SheetError(Exception):
    """The file cannot be read as a sheet. The message is meant for the person who uploaded it."""

    def __init__(self, message: str, code: str = "bad_file") -> None:
        super().__init__(message)
        self.message = message
        self.code = code


def is_blank(cells: Any) -> bool:
    return all(cell is None or (isinstance(cell, str) and not cell.strip()) for cell in cells)


class Sheet:
    """A sheet being read: the header first, then the rows one by one. Always `close()` it."""

    headers: list[str]

    def rows(self) -> Iterator[tuple[int, list[Any]]]:  # (row number as a spreadsheet shows it, the cells) - blank rows are skipped
        raise NotImplementedError

    def progress(self) -> int | None:
        """Percent read, when the size of the file tells; None for a format that does not."""
        return None

    def close(self) -> None:
        return None

    def __enter__(self) -> "Sheet":
        return self

    def __exit__(self, *exc: object) -> None:
        self.close()


# ------------------------------------------------------------------------------------------------------------------ CSV
def detect_encoding(sample: bytes) -> str:
    if sample.startswith((codecs.BOM_UTF32_LE, codecs.BOM_UTF32_BE)):
        return "utf-32"
    if sample.startswith(codecs.BOM_UTF8):
        return "utf-8-sig"
    if sample.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return "utf-16"
    if b"\x00" in sample:
        raise SheetError("This does not look like a text file (it contains binary data). Save the sheet as CSV, or upload the Excel (.xlsx) file.")
    try:
        codecs.getincrementaldecoder("utf-8")(errors="strict").decode(sample, final=False)  # (final=False: a character cut at the end is fine)
        return "utf-8"
    except UnicodeDecodeError:
        return "cp1252"  # what Excel on Windows writes for "CSV"


def sniff_delimiter(text: str) -> str:
    head = text[:8192]
    try:
        return csv.Sniffer().sniff(head, delimiters=",;\t|").delimiter
    except csv.Error:
        first = head.splitlines()[0] if head.splitlines() else ""
        counts = {d: first.count(d) for d in ",;\t|"}
        best = max(counts, key=lambda d: counts[d])
        return best if counts[best] else ","


class CsvSheet(Sheet):
    def __init__(self, path: Path, max_columns: int) -> None:
        self._max_columns = max_columns
        self._size = max(1, path.stat().st_size)
        self._file = open(path, "rb")  # noqa: SIM115 - closed in close()
        try:
            sample = self._file.read(_SAMPLE_BYTES)
            encoding = detect_encoding(sample)
            delimiter = sniff_delimiter(sample.decode(encoding, errors="replace"))
            self._file.seek(0)
            self._text = io.TextIOWrapper(self._file, encoding=encoding, errors="replace", newline="")
            self._reader = csv.reader(self._text, delimiter=delimiter)
            self._number = 0
            self.headers = self._read_headers()
        except Exception:
            self.close()
            raise

    def _next(self) -> list[str] | None:
        try:
            cells = next(self._reader)
        except StopIteration:
            return None
        except csv.Error as exc:
            raise SheetError(f"Row {self._number + 1} could not be read ({exc}). A cell may be larger than 128 KB or a quote is not closed.") from exc
        self._number += 1
        return cells

    def _read_headers(self) -> list[str]:
        while True:
            cells = self._next()
            if cells is None:
                raise SheetError("The file is empty.", code="empty_file")
            if not is_blank(cells):
                break
        headers = [cell.strip() for cell in cells]
        while headers and not headers[-1]:
            headers.pop()
        if len(headers) > self._max_columns:
            raise SheetError(f"The sheet has more than {self._max_columns} columns.")
        return headers

    def rows(self) -> Iterator[tuple[int, list[Any]]]:
        while True:
            cells = self._next()
            if cells is None:
                return
            if cells and not is_blank(cells):
                yield self._number, cells[: self._max_columns]

    def progress(self) -> int | None:
        return min(99, self._file.tell() * 100 // self._size)

    def close(self) -> None:
        try:
            self._file.close()
        except Exception:  # pragma: no cover
            pass


# ----------------------------------------------------------------------------------------------------------------- XLSX
def check_xlsx_container(path: Path) -> None:
    """Refuse a file that cannot be a sheet or that unpacks to something harmful - before anything is unpacked."""
    settings = get_settings()
    if not zipfile.is_zipfile(path):
        raise SheetError("This is not a valid Excel (.xlsx) file. Open it in Excel and save it again as .xlsx or as CSV.")
    try:
        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
            if len(infos) > _MAX_ZIP_ENTRIES:
                raise SheetError("The Excel file has an unusual structure and was not accepted.")
            names = {info.filename for info in infos}
            if "xl/workbook.xml" not in names or "[Content_Types].xml" not in names:
                raise SheetError("This is not an Excel workbook (.xlsx).")
            if any(name.lower().endswith("vbaproject.bin") for name in names):
                raise SheetError("Workbooks with macros are not accepted. Save the sheet as a plain .xlsx or as CSV.")
            if any(info.flag_bits & 0x1 for info in infos):
                raise SheetError("The Excel file is password protected. Remove the password and upload it again.")
            unpacked = sum(info.file_size for info in infos)
            if unpacked > settings.max_import_xlsx_unpacked_mb * 1024 * 1024:
                raise SheetError(f"The Excel file unpacks to more than {settings.max_import_xlsx_unpacked_mb} MB and was not accepted.")
            for info in infos:
                if info.file_size > 10 * 1024 * 1024 and info.compress_size and info.file_size / info.compress_size > _MAX_RATIO:
                    raise SheetError("The Excel file is compressed far more than a real sheet is and was not accepted.")
    except zipfile.BadZipFile as exc:
        raise SheetError("The Excel file is damaged. Open it in Excel and save it again.") from exc


class XlsxSheet(Sheet):
    def __init__(self, path: Path, max_columns: int) -> None:
        self._max_columns = max_columns
        check_xlsx_container(path)
        from openpyxl import load_workbook

        try:
            self._book = load_workbook(filename=str(path), read_only=True, data_only=True, keep_links=False)
        except Exception as exc:
            raise SheetError("The Excel file could not be read. Open it in Excel and save it again as .xlsx or as CSV.") from exc
        try:
            if not self._book.worksheets:
                raise SheetError("The workbook has no sheet.")
            self._iter = self._book.worksheets[0].iter_rows(values_only=True)  # the first sheet
            self._number = 0
            self.headers = self._read_headers()
        except Exception:
            self.close()
            raise

    def _next(self) -> tuple | None:
        try:
            cells = next(self._iter)
        except StopIteration:
            return None
        except SheetError:
            raise
        except Exception as exc:
            raise SheetError(f"The Excel file is damaged near row {self._number + 1}. Open it in Excel and save it again.") from exc
        self._number += 1
        return cells

    def _read_headers(self) -> list[str]:
        while True:
            cells = self._next()
            if cells is None:
                raise SheetError("The file is empty.", code="empty_file")
            if not is_blank(cells):
                break
        headers = ["" if cell is None else str(cell).strip() for cell in cells]
        while headers and not headers[-1]:
            headers.pop()
        if len(headers) > self._max_columns:
            raise SheetError(f"The sheet has more than {self._max_columns} columns.")
        return headers

    def rows(self) -> Iterator[tuple[int, list[Any]]]:
        while True:
            cells = self._next()
            if cells is None:
                return
            if not is_blank(cells):
                yield self._number, list(cells[: self._max_columns])

    def close(self) -> None:
        try:
            self._book.close()
        except Exception:  # pragma: no cover
            pass


def open_sheet(path: Path, file_type: str) -> Sheet:
    max_columns = get_settings().max_import_columns
    return XlsxSheet(path, max_columns) if file_type == "xlsx" else CsvSheet(path, max_columns)
