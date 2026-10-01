"""Preview-first, transactional contact import (CSV / XLSX) - section 15 of the documentation.

Flow:  upload -> validate in background (status `previewed`) -> admin reviews counts ->
       confirm -> apply in ONE database transaction (status `completed`).
"""

from __future__ import annotations

import csv
import io
import logging
import os
import re
import uuid
from collections import defaultdict
from collections.abc import Iterable, Iterator, Sequence
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from fastapi import Request, UploadFile
from sqlalchemy import func, insert, select, update
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.core.database import new_session
from app.core.errors import Conflict, NotFound, PayloadTooLarge, ValidationFailed
from app.core.redis_client import redis_lock
from app.core.timeutils import utcnow
from app.models.contact import Campaign, Contact, ContactAssignment
from app.models.employee import Employee
from app.models.imports import (
    IMPORT_APPLYING,
    IMPORT_CANCELLED,
    IMPORT_COMPLETED,
    IMPORT_FAILED,
    IMPORT_PREVIEWED,
    IMPORT_VALIDATING,
    ROW_DUPLICATE,
    ROW_INVALID,
    ROW_VALID,
    Import,
    ImportRow,
)
from app.schemas.common import validate_email
from app.services import assignment_service, audit_service
from app.services.contact_service import build_search_text, refresh_search_text
from app.services.notification_service import notify
from app.services.phone import clean_phone_input, normalize_phone

log = logging.getLogger(__name__)

ALLOWED_EXTENSIONS = {".csv", ".xlsx"}
CHUNK = 1000

HEADER_ALIASES: dict[str, set[str]] = {
    "name": {"name", "full name", "fullname", "customer name", "contact name", "client name", "lead name", "naam"},
    "phone": {
        "phone", "mobile", "mobile number", "mobile no", "mobile no.", "phone number", "phone no", "contact number",
        "contact no", "contact", "whatsapp", "whatsapp number", "cell", "cell number", "number", "mobile1",
    },
    "email": {"email", "email id", "e mail", "mail", "email address"},
    "location": {"location", "city", "area", "address", "place", "town", "district"},
    "category": {"category", "segment", "type", "customer type", "lead type"},
    "priority": {"priority"},
    "tags": {"tags", "tag", "labels", "label"},
    "assigned_to": {"assigned to", "assignee", "employee", "employee code", "employee id", "agent", "assigned employee"},
}
_ALIAS_LOOKUP = {alias: canonical for canonical, aliases in HEADER_ALIASES.items() for alias in aliases}

_PRIORITY_WORDS = {
    "1": 1, "high": 1, "h": 1, "urgent": 1,
    "2": 2, "medium": 2, "med": 2, "m": 2, "normal": 2, "": 2,
    "3": 3, "low": 3, "l": 3,
}


def _norm_header(value: Any) -> str:
    return re.sub(r"[\s_\-]+", " ", str(value or "").strip().lower())


def _chunks(items: Sequence, size: int = CHUNK) -> Iterable[Sequence]:
    for i in range(0, len(items), size):
        yield items[i : i + size]


# ---------------------------------------------------------------- file reading
def _read_csv(path: Path) -> tuple[list[str], list[list[Any]]]:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover - latin-1 never fails
        text = raw.decode("latin-1")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows = [r for r in reader]
    if not rows:
        return [], []
    return [str(h) for h in rows[0]], rows[1:]


def _read_xlsx(path: Path) -> tuple[list[str], list[list[Any]]]:
    from openpyxl import load_workbook

    workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        iterator = sheet.iter_rows(values_only=True)
        try:
            header = next(iterator)
        except StopIteration:
            return [], []
        headers = ["" if h is None else str(h) for h in header]
        rows = [list(r) for r in iterator]
        return headers, rows
    finally:
        workbook.close()


def read_table(path: Path, file_type: str) -> tuple[list[str], list[list[Any]]]:
    headers, rows = _read_xlsx(path) if file_type == "xlsx" else _read_csv(path)
    # drop fully empty trailing/blank rows
    rows = [r for r in rows if any(str(c).strip() for c in r if c is not None)]
    return headers, rows


# ------------------------------------------------------------------ validation
@dataclass
class RowResult:
    row_number: int
    status: str
    data: dict[str, Any]
    errors: list[str] = field(default_factory=list)
    normalized_phone: str | None = None
    duplicate_of: str | None = None


def map_headers(headers: list[str]) -> tuple[dict[int, str], dict[int, str]]:
    canonical: dict[int, str] = {}
    custom: dict[int, str] = {}
    used: set[str] = set()
    for idx, header in enumerate(headers):
        key = _ALIAS_LOOKUP.get(_norm_header(header))
        if key and key not in used:
            canonical[idx] = key
            used.add(key)
        elif str(header).strip():
            custom[idx] = str(header).strip()[:60]
    return canonical, custom


def _cell(row: list[Any], idx: int) -> Any:
    return row[idx] if idx < len(row) else None


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def validate_row(
    row_number: int,
    row: list[Any],
    canonical: dict[int, str],
    custom: dict[int, str],
    *,
    default_priority: int,
    employees: dict[str, int],
) -> RowResult:
    values = {key: _cell(row, idx) for idx, key in canonical.items()}
    errors: list[str] = []

    name = _text(values.get("name"))
    if not name:
        errors.append("Name is required.")
    elif len(name) > 255:
        errors.append("Name is longer than 255 characters.")

    raw_phone = clean_phone_input(values.get("phone"))
    normalized = None
    if not raw_phone:
        errors.append("Mobile number is required.")
    else:
        normalized = normalize_phone(raw_phone)
        if normalized is None:
            errors.append(f"Invalid mobile number '{raw_phone[:32]}'.")

    email = _text(values.get("email")) or None
    if email:
        try:
            email = validate_email(email)
        except ValueError:
            errors.append(f"Invalid email '{email[:60]}'.")
            email = None

    priority = default_priority
    prio_text = _text(values.get("priority")).lower()
    if prio_text:
        if prio_text in _PRIORITY_WORDS:
            priority = _PRIORITY_WORDS[prio_text]
        else:
            errors.append(f"Invalid priority '{prio_text[:20]}' (use 1-3 or High/Medium/Low).")

    tags = [t.strip()[:40] for t in re.split(r"[;,|]", _text(values.get("tags"))) if t.strip()][:20]

    custom_fields: dict[str, Any] = {}
    for idx, header in custom.items():
        cell = _text(_cell(row, idx))
        if cell:
            custom_fields[header] = cell[:500]
        if len(custom_fields) >= 30:
            break

    assigned_employee_id = None
    assigned_key = _text(values.get("assigned_to"))
    if assigned_key:
        assigned_employee_id = employees.get(assigned_key.lower())
        if assigned_employee_id is None:
            errors.append(f"Unknown or inactive employee '{assigned_key[:40]}'.")

    if errors:
        raw = {(custom.get(i) or canonical.get(i) or f"col{i}"): _text(_cell(row, i)) for i in range(len(row))}
        return RowResult(row_number, ROW_INVALID, raw, errors, normalized)

    data = {
        "name": name,
        "phone_raw": raw_phone[:64],
        "email": email,
        "location": _text(values.get("location"))[:255] or None,
        "category": _text(values.get("category"))[:100] or None,
        "priority": priority,
        "tags": tags,
        "custom_fields": custom_fields,
        "assigned_employee_id": assigned_employee_id,
    }
    return RowResult(row_number, ROW_VALID, data, [], normalized)


# ------------------------------------------------------------- import lifecycle
def _storage_dir() -> Path:
    path = Path(get_settings().import_storage_path)
    path.mkdir(parents=True, exist_ok=True)
    return path


def _remove_file(imp: Import) -> None:
    try:
        Path(imp.file_path).unlink(missing_ok=True)
    except OSError:  # pragma: no cover
        log.warning("Could not delete import file %s", imp.file_path)


def save_upload(upload: UploadFile) -> tuple[Path, str, int]:
    settings = get_settings()
    filename = os.path.basename(upload.filename or "upload")
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise ValidationFailed("Only .csv and .xlsx files can be imported.", code="unsupported_file_type")
    target = _storage_dir() / f"{uuid.uuid4().hex}{ext}"
    limit = settings.max_import_mb * 1024 * 1024
    size = 0
    try:
        with open(target, "wb") as out:
            while True:
                chunk = upload.file.read(1024 * 1024)
                if not chunk:
                    break
                size += len(chunk)
                if size > limit:
                    raise PayloadTooLarge(f"File is larger than {settings.max_import_mb} MB.", code="file_too_large")
                out.write(chunk)
    except Exception:
        target.unlink(missing_ok=True)
        raise
    if size == 0:
        target.unlink(missing_ok=True)
        raise ValidationFailed("The uploaded file is empty.", code="empty_file")
    return target, ext.lstrip("."), size


def create_import(
    db: Session,
    *,
    upload: UploadFile,
    actor: Employee,
    mode: str,
    campaign_id: int | None,
    assign_employee_ids: list[int],
    assign_strategy: str,
    default_priority: int,
    request: Request,
) -> Import:
    if mode not in ("skip", "update"):
        raise ValidationFailed("mode must be 'skip' or 'update'.", code="bad_mode")
    if assign_strategy not in ("round_robin", "balanced"):
        raise ValidationFailed("assign_strategy must be 'round_robin' or 'balanced'.", code="bad_strategy")
    if default_priority not in (1, 2, 3):
        raise ValidationFailed("default_priority must be 1, 2 or 3.", code="bad_priority")
    if campaign_id is not None and db.get(Campaign, campaign_id) is None:
        raise ValidationFailed("Campaign does not exist.", code="unknown_campaign")
    if assign_employee_ids:
        assignment_service.resolve_target_employees(db, assign_employee_ids, None)

    path, file_type, size = save_upload(upload)
    imp = Import(
        created_by=actor.id,
        filename=os.path.basename(upload.filename or "upload")[:255],
        file_type=file_type,
        file_path=str(path),
        file_size=size,
        mode=mode,
        status=IMPORT_VALIDATING,
        campaign_id=campaign_id,
        options={
            "assign_employee_ids": assign_employee_ids,
            "assign_strategy": assign_strategy,
            "default_priority": default_priority,
        },
    )
    db.add(imp)
    db.flush()
    audit_service.record(
        db, action="import.upload", actor=actor, entity_type="import", entity_id=imp.id, request=request,
        details={"filename": imp.filename, "size": size, "mode": mode},
    )
    db.commit()
    return imp


def _fail(db: Session, imp: Import, message: str) -> None:
    db.rollback()
    imp = db.get(Import, imp.id)  # type: ignore[assignment]
    imp.status = IMPORT_FAILED
    imp.error_message = message[:2000]
    db.commit()
    _remove_file(imp)


def run_validation(import_id: int) -> None:
    """Background task: parse and validate the file, persist row results, set status `previewed`."""
    db = new_session()
    try:
        imp = db.get(Import, import_id)
        if imp is None or imp.status != IMPORT_VALIDATING:
            return
        settings = get_settings()
        try:
            headers, rows = read_table(Path(imp.file_path), imp.file_type)
        except Exception as exc:  # corrupt file
            log.warning("Import %s unreadable: %s", import_id, exc)
            _fail(db, imp, "The file could not be read. Make sure it is a valid CSV or XLSX file.")
            return
        if not headers:
            _fail(db, imp, "The file has no header row.")
            return
        if len(rows) > settings.max_import_rows:
            _fail(db, imp, f"The file has more than {settings.max_import_rows} rows. Split it into smaller files.")
            return

        canonical, custom = map_headers(headers)
        if "phone" not in canonical.values():
            _fail(db, imp, "No mobile number column found. Name a column 'Mobile' or 'Phone'.")
            return
        if "name" not in canonical.values():
            _fail(db, imp, "No name column found. Name a column 'Name'.")
            return

        employees: dict[str, int] = {}
        for emp_id, code, email in db.execute(select(Employee.id, Employee.employee_code, Employee.email).where(Employee.is_active.is_(True))):
            employees[code.lower()] = emp_id
            employees[email.lower()] = emp_id

        default_priority = int((imp.options or {}).get("default_priority", 2))
        results = [
            validate_row(i, row, canonical, custom, default_priority=default_priority, employees=employees)
            for i, row in enumerate(rows, start=2)  # row 1 is the header
        ]

        # duplicates inside the file: the first valid occurrence wins
        seen: dict[str, int] = {}
        for res in results:
            if res.status != ROW_VALID or not res.normalized_phone:
                continue
            if res.normalized_phone in seen:
                res.status = ROW_DUPLICATE
                res.duplicate_of = "file"
                res.errors = [f"Same mobile number as row {seen[res.normalized_phone]}."]
            else:
                seen[res.normalized_phone] = res.row_number

        # duplicates against existing contacts
        phones = list(seen.keys())
        existing: dict[str, tuple[int, bool]] = {}
        for chunk in _chunks(phones):
            for cid, phone, deleted_at in db.execute(
                select(Contact.id, Contact.normalized_phone, Contact.deleted_at).where(Contact.normalized_phone.in_(chunk))
            ):
                existing[phone] = (cid, deleted_at is not None)
        for res in results:
            if res.status == ROW_VALID and res.normalized_phone in existing:
                cid, deleted = existing[res.normalized_phone]
                if deleted:
                    res.data["restore_contact_id"] = cid  # a previously deleted contact is revived
                else:
                    res.status = ROW_DUPLICATE
                    res.duplicate_of = "existing"
                    res.errors = ["Mobile number already exists in contacts."]
                    res.data["existing_contact_id"] = cid

        counts = {ROW_VALID: 0, ROW_INVALID: 0, ROW_DUPLICATE: 0}
        payload = []
        for res in results:
            counts[res.status] += 1
            payload.append(
                {
                    "import_id": imp.id,
                    "row_number": res.row_number,
                    "status": res.status,
                    "duplicate_of": res.duplicate_of,
                    "action": None,
                    "normalized_phone": res.normalized_phone,
                    "data": res.data,
                    "errors": res.errors or None,
                }
            )
        for chunk in _chunks(payload, 500):
            db.execute(insert(ImportRow), list(chunk))

        imp.total_rows = len(results)
        imp.valid_rows = counts[ROW_VALID]
        imp.invalid_rows = counts[ROW_INVALID]
        imp.duplicate_rows = counts[ROW_DUPLICATE]
        imp.status = IMPORT_PREVIEWED
        imp.error_message = None
        db.commit()
    except Exception as exc:  # pragma: no cover - defensive
        log.exception("Import validation crashed")
        try:
            imp = db.get(Import, import_id)
            if imp is not None:
                _fail(db, imp, f"Validation failed: {exc}")
        except Exception:
            log.exception("Could not mark import %s as failed", import_id)
    finally:
        db.close()


def get_import(db: Session, import_id: int) -> Import:
    imp = db.get(Import, import_id)
    if imp is None:
        raise NotFound("Import not found.")
    return imp


def list_imports(db: Session, *, page: int, page_size: int) -> tuple[list[Import], int]:
    total = db.scalar(select(func.count(Import.id))) or 0
    rows = db.scalars(select(Import).order_by(Import.created_at.desc(), Import.id.desc()).limit(page_size).offset((page - 1) * page_size)).all()
    return list(rows), total


def list_rows(db: Session, import_id: int, *, status: str | None, page: int, page_size: int) -> tuple[list[ImportRow], int]:
    stmt = select(ImportRow).where(ImportRow.import_id == import_id)
    if status:
        stmt = stmt.where(ImportRow.status == status)
    total = db.scalar(select(func.count()).select_from(stmt.subquery())) or 0
    rows = db.scalars(stmt.order_by(ImportRow.row_number).limit(page_size).offset((page - 1) * page_size)).all()
    return list(rows), total


def cancel_import(db: Session, imp: Import, actor: Employee, request: Request) -> Import:
    if imp.status not in (IMPORT_PREVIEWED, IMPORT_VALIDATING, IMPORT_FAILED):
        raise Conflict(f"An import in status '{imp.status}' cannot be cancelled.", code="bad_import_state")
    imp.status = IMPORT_CANCELLED
    _remove_file(imp)
    audit_service.record(db, action="import.cancel", actor=actor, entity_type="import", entity_id=imp.id, request=request)
    db.commit()
    return imp


def begin_apply(db: Session, imp: Import, actor: Employee, mode: str | None, request: Request) -> Import:
    if mode is not None and mode not in ("skip", "update"):
        raise ValidationFailed("mode must be 'skip' or 'update'.", code="bad_mode")
    if imp.status != IMPORT_PREVIEWED:
        raise Conflict(f"Only a previewed import can be confirmed (current status: {imp.status}).", code="bad_import_state")
    claimed = db.execute(
        update(Import)
        .where(Import.id == imp.id, Import.status == IMPORT_PREVIEWED)
        .values(status=IMPORT_APPLYING, confirmed_at=utcnow(), mode=mode or imp.mode)
    )
    if not claimed.rowcount:
        db.rollback()
        raise Conflict("This import is already being processed.", code="bad_import_state")
    audit_service.record(
        db, action="import.confirm", actor=actor, entity_type="import", entity_id=imp.id, request=request, details={"mode": mode or imp.mode}
    )
    db.commit()
    db.refresh(imp)
    return imp


def _iter_rows(db: Session, import_id: int, status: str) -> Iterator[list[ImportRow]]:
    last_id = 0
    while True:
        rows = db.scalars(
            select(ImportRow)
            .where(ImportRow.import_id == import_id, ImportRow.status == status, ImportRow.id > last_id)
            .order_by(ImportRow.id)
            .limit(CHUNK)
        ).all()
        if not rows:
            return
        last_id = rows[-1].id
        yield list(rows)


def _apply_fields(contact: Contact, data: dict[str, Any], *, reset: bool) -> None:
    contact.name = data["name"]
    contact.phone_raw = data["phone_raw"]
    for key in ("email", "location", "category"):
        if data.get(key):
            setattr(contact, key, data[key])
    if data.get("priority"):
        contact.priority = data["priority"]
    if data.get("tags"):
        merged = list(contact.tags or [])
        for tag in data["tags"]:
            if tag.lower() not in (t.lower() for t in merged):
                merged.append(tag)
        contact.tags = merged[:20]
    if data.get("custom_fields"):
        merged_cf = dict(contact.custom_fields or {})
        merged_cf.update(data["custom_fields"])
        contact.custom_fields = merged_cf
    if reset:
        contact.deleted_at = None
        contact.status = "new"
        contact.failed_attempts = 0
        contact.next_eligible_at = None
    refresh_search_text(contact)


def run_apply(import_id: int, actor_id: int | None) -> None:
    """Background task: apply a previewed import inside a single transaction."""
    db = new_session()
    try:
        imp = db.get(Import, import_id)
        if imp is None or imp.status != IMPORT_APPLYING:
            return
        with redis_lock(f"import-apply:{import_id}", ttl_seconds=900) as acquired:
            if not acquired:
                return
            try:
                _apply(db, imp, actor_id)
            except Exception as exc:
                log.exception("Import %s failed while applying; rolled back", import_id)
                _fail(db, imp, f"Import failed and was rolled back: {exc}")
    finally:
        db.close()


def _apply(db: Session, imp: Import, actor_id: int | None) -> None:
    now = utcnow()
    mode = imp.mode
    options = imp.options or {}
    inserted = updated = skipped = 0
    touched_contact_ids: list[int] = []  # contacts created/updated by this import

    # 1) new contacts (bulk insert) and revived contacts
    for rows in _iter_rows(db, imp.id, ROW_VALID):
        new_rows = [r for r in rows if not r.data.get("restore_contact_id")]
        restore_rows = [r for r in rows if r.data.get("restore_contact_id")]
        if new_rows:
            db.execute(
                insert(Contact),
                [
                    {
                        "name": r.data["name"],
                        "phone_raw": r.data["phone_raw"],
                        "normalized_phone": r.normalized_phone,
                        "email": r.data.get("email"),
                        "location": r.data.get("location"),
                        "category": r.data.get("category"),
                        "priority": r.data.get("priority", 2),
                        "tags": r.data.get("tags") or [],
                        "custom_fields": r.data.get("custom_fields") or {},
                        "search_text": build_search_text(
                            name=r.data["name"],
                            phone_raw=r.data["phone_raw"],
                            normalized_phone=r.normalized_phone or "",
                            email=r.data.get("email"),
                            location=r.data.get("location"),
                            category=r.data.get("category"),
                            tags=r.data.get("tags") or [],
                            custom_fields=r.data.get("custom_fields") or {},
                        ),
                        "status": "new",
                        "source": "import",
                        "import_id": imp.id,
                        "call_count": 0,
                        "failed_attempts": 0,
                        "created_by": actor_id,
                        "created_at": now,
                        "updated_at": now,
                    }
                    for r in new_rows
                ],
            )
            inserted += len(new_rows)
        if restore_rows:
            contacts = {c.id: c for c in db.scalars(select(Contact).where(Contact.id.in_([r.data["restore_contact_id"] for r in restore_rows])))}
            for r in restore_rows:
                contact = contacts.get(r.data["restore_contact_id"])
                if contact is not None:
                    _apply_fields(contact, r.data, reset=True)
                    contact.import_id = imp.id
                    contact.source = "import"
                    inserted += 1
        db.flush()
        db.execute(update(ImportRow).where(ImportRow.id.in_([r.id for r in rows])).values(action="insert"))

    # 2) duplicates
    for rows in _iter_rows(db, imp.id, ROW_DUPLICATE):
        if mode == "update":
            existing_rows = [r for r in rows if r.duplicate_of == "existing"]
            in_file_rows = [r for r in rows if r.duplicate_of != "existing"]
            contacts = {
                c.id: c
                for c in db.scalars(select(Contact).where(Contact.id.in_([r.data["existing_contact_id"] for r in existing_rows])))
            }
            for r in existing_rows:
                contact = contacts.get(r.data["existing_contact_id"])
                if contact is not None:
                    _apply_fields(contact, r.data, reset=False)
                    updated += 1
                else:
                    skipped += 1
            skipped += len(in_file_rows)
            db.flush()
            db.execute(update(ImportRow).where(ImportRow.id.in_([r.id for r in existing_rows])).values(action="update"))
            if in_file_rows:
                db.execute(update(ImportRow).where(ImportRow.id.in_([r.id for r in in_file_rows])).values(action="skip"))
        else:
            skipped += len(rows)
            db.execute(update(ImportRow).where(ImportRow.id.in_([r.id for r in rows])).values(action="skip"))

    # 3) link import rows to contacts (one correlated UPDATE, uses the unique phone index)
    db.execute(
        update(ImportRow)
        .where(ImportRow.import_id == imp.id, ImportRow.action.in_(("insert", "update")))
        .values(contact_id=select(Contact.id).where(Contact.normalized_phone == ImportRow.normalized_phone).scalar_subquery())
    )

    touched_contact_ids = list(
        db.scalars(select(ImportRow.contact_id).where(ImportRow.import_id == imp.id, ImportRow.contact_id.is_not(None)))
    )

    # 4) campaign membership
    if imp.campaign_id is not None and touched_contact_ids:
        assignment_service.attach_to_campaign(db, imp.campaign_id, touched_contact_ids)

    # 5) assignment rules: per-row `assigned_to` first, then distribute the remainder
    assigned = 0
    if touched_contact_ids:
        rows_for_assign = db.execute(
            select(ImportRow.contact_id, ImportRow.data).where(
                ImportRow.import_id == imp.id, ImportRow.contact_id.is_not(None), ImportRow.action.in_(("insert", "update"))
            )
        ).all()
        already_assigned: set[int] = set()
        for chunk in _chunks([r[0] for r in rows_for_assign]):
            already_assigned |= set(
                db.scalars(select(ContactAssignment.contact_id).where(ContactAssignment.contact_id.in_(chunk), ContactAssignment.status == "active"))
            )

        explicit: dict[int, list[int]] = defaultdict(list)
        remainder: list[int] = []
        for contact_id, data in rows_for_assign:
            if contact_id in already_assigned:
                continue
            employee_id = (data or {}).get("assigned_employee_id")
            if employee_id:
                explicit[employee_id].append(contact_id)
            else:
                remainder.append(contact_id)

        plan: dict[int, list[int]] = defaultdict(list)
        for eid, cids in explicit.items():
            plan[eid].extend(cids)
        pool_ids = options.get("assign_employee_ids") or []
        if remainder and pool_ids:
            employees = assignment_service.resolve_target_employees(db, pool_ids, None)
            loads = assignment_service.current_loads(db, [e.id for e in employees])
            for eid, cids in assignment_service.plan_distribution(remainder, employees, options.get("assign_strategy", "round_robin"), loads).items():
                plan[eid].extend(cids)

        if plan:
            assigned, _ = assignment_service.apply_plan(
                db, plan=dict(plan), campaign_id=imp.campaign_id, actor_id=actor_id, current={}
            )
            for eid, cids in plan.items():
                notify(
                    db, eid, type="assignment",
                    title=f"{len(cids)} new contact{'s' if len(cids) != 1 else ''} assigned",
                    body="Open your queue to start calling.", data={"count": len(cids), "import_id": imp.id},
                )

    imp.inserted_rows = inserted
    imp.updated_rows = updated
    imp.skipped_rows = skipped
    imp.assigned_rows = assigned
    imp.status = IMPORT_COMPLETED
    imp.completed_at = utcnow()
    imp.error_message = None
    audit_service.record(
        db,
        action="import.completed",
        actor=db.get(Employee, actor_id) if actor_id else None,
        entity_type="import",
        entity_id=imp.id,
        details={"inserted": inserted, "updated": updated, "skipped": skipped, "assigned": assigned, "invalid": imp.invalid_rows},
    )
    db.commit()
    _remove_file(imp)


# ---------------------------------------------------------------- error report
def _csv_safe(value: Any) -> str:
    text = "" if value is None else str(value)
    if text[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + text  # neutralise spreadsheet formula injection
    return text


def export_issues_csv(db: Session, imp: Import) -> str:
    out = io.StringIO()
    writer = csv.writer(out)
    rows = db.scalars(
        select(ImportRow).where(ImportRow.import_id == imp.id, ImportRow.status.in_((ROW_INVALID, ROW_DUPLICATE))).order_by(ImportRow.row_number)
    ).all()
    columns: list[str] = []
    for r in rows:
        for key in (r.data or {}):
            if key not in columns and key not in ("assigned_employee_id", "existing_contact_id", "restore_contact_id"):
                columns.append(key)
    writer.writerow(["row_number", "status", "problem", *columns])
    for r in rows:
        writer.writerow(
            [r.row_number, r.status, _csv_safe("; ".join(r.errors or [])), *[_csv_safe((r.data or {}).get(c)) for c in columns]]
        )
    return out.getvalue()


def purge_old_imports(db: Session, *, older_than_days: int = 30) -> int:
    """Retention: drop row-level import data (contains personal data) after N days."""
    from datetime import timedelta

    cutoff = utcnow() - timedelta(days=older_than_days)
    old_ids = list(db.scalars(select(Import.id).where(Import.created_at < cutoff)))
    if not old_ids:
        return 0
    db.execute(ImportRow.__table__.delete().where(ImportRow.import_id.in_(old_ids)))
    db.commit()
    return len(old_ids)
